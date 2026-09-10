#!/usr/bin/env python3
"""
Render a .pptx to PNGs by reading the file back and drawing it in Chrome.

    python3 tools/carousel/preview_pptx.py out/carousel/<deck>.pptx

Why this exists: there is no pptx renderer on this Mac. LibreOffice is not
installed, Keynote will not open a file from AppleScript here, and QuickLook
has no generator for pptx. Without something like this, a generated deck would
ship having never been looked at.

It deliberately reads the SAVED FILE rather than the builder's intent - shape
positions, fills, fonts and colours all come back out of the XML - so what you
see is what is actually in the deck. It approximates: Chrome's line breaking is
not PowerPoint's, and it ignores anything the builder does not use.
"""

import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
PX_IN = 108                      # 1080px / 10in
PT_PX = PX_IN / 72.0


def px(emu):
    return (emu or 0) / EMU_IN * PX_IN


def rgb_of(obj):
    """RGBColor subclasses tuple, so "#%s" % color raises "not all arguments
    converted" and a bare except would turn every colour into None. Format it
    with f-string interpolation instead."""
    try:
        return f"#{obj.rgb}"
    except Exception:                                          # noqa: BLE001
        return None


def shape_html(sh):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    css = (f"left:{px(sh.left):.1f}px;top:{px(sh.top):.1f}px;"
           f"width:{px(sh.width):.1f}px;height:{px(sh.height):.1f}px;")

    # connectors carry begin/end rather than a box
    if sh.shape_type == MSO_SHAPE_TYPE.LINE:
        x1, y1 = px(sh.begin_x), px(sh.begin_y)
        x2, y2 = px(sh.end_x), px(sh.end_y)
        col = rgb_of(sh.line.color) or "#000"
        w = max((sh.line.width or 9525) / EMU_IN * PX_IN, 1)
        return (f'<div style="position:absolute;left:{min(x1,x2):.1f}px;'
                f'top:{min(y1,y2):.1f}px;width:{max(abs(x2-x1),w):.1f}px;'
                f'height:{max(abs(y2-y1),w):.1f}px;background:{col}"></div>')

    box = ""
    fill = None
    try:
        if sh.fill.type is not None and sh.fill.type == 1:      # MSO_FILL.SOLID
            fill = rgb_of(sh.fill.fore_color)
    except Exception:                                          # noqa: BLE001
        pass
    if fill:
        box += f"background:{fill};"
    try:
        lc = rgb_of(sh.line.color)
        if lc and sh.line.width:
            box += f"border:{(sh.line.width/EMU_IN*PX_IN):.1f}px solid {lc};"
    except Exception:                                          # noqa: BLE001
        pass

    inner = ""
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            # PP_ALIGN: LEFT=1, CENTER=2, RIGHT=3.
            align = {2: "center", 3: "right"}.get(
                p.alignment.value if p.alignment else None, "left")
            ls = p.line_spacing or 1.2
            runs = ""
            for r in p.runs:
                f = r.font
                size = (f.size.pt if f.size else 18) * PT_PX
                col = rgb_of(f.color) or "#000"
                spc = r.font._rPr.get("spc")
                tr = f"letter-spacing:{int(spc)/100*PT_PX:.2f}px;" if spc else ""
                runs += (f'<span style="font-family:\'{f.name}\';'
                         f'font-size:{size:.1f}px;color:{col};'
                         f'{"font-weight:700;" if f.bold else ""}{tr}">'
                         f'{r.text or "&nbsp;"}</span>')
            inner += f'<p style="margin:0;text-align:{align};line-height:{ls}">{runs}</p>'
    return f'<div style="position:absolute;{css}{box}">{inner}</div>'


def main(path):
    prs = Presentation(path)
    W = px(prs.slide_width)
    H = px(prs.slide_height)
    out = Path(path).with_suffix("")
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        for i, slide in enumerate(prs.slides, 1):
            bg = rgb_of(slide.background.fill.fore_color) or "#FFF"
            body = "".join(shape_html(sh) for sh in slide.shapes)
            html = (f'<!doctype html><meta charset="utf-8"><style>'
                    f'body{{margin:0}}div.slide{{position:relative;'
                    f'width:{W:.0f}px;height:{H:.0f}px;background:{bg};'
                    f'overflow:hidden}}</style>'
                    f'<div class="slide">{body}</div>')
            f = Path(tmp) / f"s{i}.html"
            f.write_text(html, encoding="utf-8")
            png = out / f"preview-{i:02d}.png"
            subprocess.run([
                "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
                "--headless", "--disable-gpu", "--hide-scrollbars",
                f"--screenshot={png}", f"--window-size={int(W)},{int(H)}",
                "--virtual-time-budget=1200", "--no-sandbox", f.as_uri(),
            ], check=True, capture_output=True)
            print(f"  {png.name}")
    print(f"-> {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1
         else "out/carousel/AI-Profit-Lab-Carousel-Template.pptx")
