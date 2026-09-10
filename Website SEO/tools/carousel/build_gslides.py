#!/usr/bin/env python3
"""
AI Profit Lab - Instagram carousel template as an editable presentation.

Writes a .pptx that Google Drive converts to a native Google Slides deck, so
the carousel can be edited in the browser instead of in JSON.

    python3 tools/carousel/build_gslides.py            # -> out/carousel/AI-Profit-Lab-Carousel-Template.pptx

WHAT IS DIFFERENT FROM THE PNG TEMPLATE
---------------------------------------
Minimal chrome, per Nahid: no wordmark, no series/counter header, no swipe
cue. What is left is the ground, four corner ticks, the type, and the single
amber accent. Everything that used to be furniture is now speaker notes.

FONTS
-----
Titles are Freehand. Body is Tenor Sans, standing in for Optima: Optima is
installed on this Mac but is NOT in the Google Slides font picker, which only
offers Google Fonts, so an Optima-set deck silently renders as Arial in the
browser. Tenor Sans is the closest humanist face that Slides can actually
serve. The local PNG builder keeps real Optima - see README.

Both faces are cached as TTFs in tools/carousel/fonts/ and every size on every
slide is measured against those real metrics (PIL) and shrunk until it fits its
box. Freehand is ~26% narrower than Tenor Sans at the same size, so no single
hardcoded scale would have worked for both.
"""

import subprocess
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
FONTS = HERE / "fonts"
OUT = REPO / "out" / "carousel"

# 4:5, the tallest ratio Instagram renders in feed. 10in keeps the numbers
# round; Slides exports 1600px on the long edge, which downsamples cleanly.
PAGE_W, PAGE_H = 10.0, 12.5
MARGIN = 0.9
CONTENT_W = PAGE_W - 2 * MARGIN

TITLE_FONT, BODY_FONT = "Freehand", "Tenor Sans"
TITLE_TTF, BODY_TTF = FONTS / "Freehand.ttf", FONTS / "TenorSans.ttf"

LIGHT = {"bg": "F1EFE8", "text": "232B26", "muted": "5A665D",
         "line": "DED8C8", "amber": "BA7517", "kicker": "5A665D"}
# --line-dark is cream at 16% over the teal; a pptx fill cannot carry alpha,
# so it is pre-blended here.
DARK = {"bg": "0A3D30", "text": "F1EFE8", "muted": "E8C98F",
        "line": "2F5A4D", "amber": "D89234", "kicker": "E8C98F"}

PT_IN = 72.0


def rgb(h):
    return RGBColor.from_string(h)


# --------------------------------------------------------------------------
# Measurement. Everything below sizes itself against the real font files.
# --------------------------------------------------------------------------
def wrap(text, ttf, pt, max_w_in):
    """Greedy wrap at a given point size. Widths come from the actual TTF."""
    font = ImageFont.truetype(str(ttf), max(int(pt), 1))
    limit = max_w_in * PT_IN
    lines, line = [], ""
    for word in text.split():
        trial = f"{line} {word}".strip()
        if font.getlength(trial) <= limit or not line:
            line = trial
        else:
            lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def fit(text, ttf, pt, max_w_in, max_lines, floor=12):
    """Shrink until the copy fits in max_lines. Returns (size, lines)."""
    size = pt
    while size > floor:
        lines = wrap(text, ttf, size, max_w_in)
        if len(lines) <= max_lines:
            return size, lines
        size -= 2
    return size, wrap(text, ttf, size, max_w_in)


def block_h(lines, pt, leading):
    return len(lines) * pt * leading / PT_IN          # inches


# --------------------------------------------------------------------------
# Drawing
# --------------------------------------------------------------------------
def textbox(slide, x, y, w, h, text, *, font, size, color, leading=1.3,
            bold=False, tracking=None, caps=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, para in enumerate((text or "").split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = leading
        run = p.add_run()
        run.text = para.upper() if caps else para
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        if tracking:
            # Letter-spacing has no python-pptx API; spc on rPr is the
            # underlying attribute, in hundredths of a point.
            run.font._rPr.set("spc", str(int(tracking * 100)))
    return box


def hairline(slide, x, y, w, color, pt=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                   Inches(x + w), Inches(y))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(pt)
    return c


def corner_ticks(slide, pal):
    """The blueprint grammar, reduced to four marks and no words."""
    inset, leg = 0.42, 0.34
    for cx, cy, dx, dy in ((inset, inset, 1, 1), (PAGE_W - inset, inset, -1, 1),
                           (PAGE_W - inset, PAGE_H - inset, -1, -1),
                           (inset, PAGE_H - inset, 1, -1)):
        for x2, y2 in ((cx + leg * dx, cy), (cx, cy + leg * dy)):
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx),
                                           Inches(cy), Inches(x2), Inches(y2))
            c.line.color.rgb = rgb(pal["line"])
            c.line.width = Pt(0.75)


def new_slide(prs, pal, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])      # blank
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(pal["bg"])
    corner_ticks(slide, pal)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def stack(items, top=None):
    """Vertically centre a list of (height, draw_fn) blocks in the page."""
    total = sum(h for h, _ in items)
    y = top if top is not None else (PAGE_H - total) / 2
    for h, draw in items:
        draw(y)
        y += h


# --------------------------------------------------------------------------
# Slides
# --------------------------------------------------------------------------
KICKER = dict(font=BODY_FONT, size=13, leading=1.2, tracking=2.2, caps=True)


def s_cover(prs):
    pal = DARK
    s = new_slide(prs, pal, notes=(
        "COVER — the whole post's hook.\n"
        "One claim, stated flat. No hype. The amber bar is this slide's single "
        "accent, so do not colour a word as well.\n"
        "Keep the title under ~60 characters or it drops a size."))
    kicker = "«For whom — e.g. Oman distributors»"
    title = "«The hook — one claim, stated flat.»"
    body = "«One line of set-up: what they are about to get.»"
    ts, tl = fit(title, TITLE_TTF, 62, CONTENT_W, 3)
    bs, bl = fit(body, BODY_TTF, 21, CONTENT_W, 3)
    stack([
        (0.42, lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.3, kicker,
                                 color=pal["kicker"], **KICKER)),
        (0.34, lambda y: _bar(s, y, pal)),
        (block_h(tl, ts, 1.16) + 0.34,
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 2.6, title, font=TITLE_FONT,
                           size=ts, color=pal["text"], leading=1.16)),
        (block_h(bl, bs, 1.45),
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 1.2, body, font=BODY_FONT,
                           size=bs, color=pal["text"], leading=1.45)),
    ])


def _bar(s, y, pal):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(y),
                             Inches(0.95), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(pal["amber"])
    bar.line.fill.background()
    bar.shadow.inherit = False


def s_text(prs):
    pal = LIGHT
    s = new_slide(prs, pal, notes=(
        "PROBLEM — name the symptom they recognise in their own week.\n"
        "Describe it the way they would describe it to a friend. Do not sell "
        "yet. The small line at the bottom is an optional margin note; delete "
        "the box if you don't need it."))
    title = "«The symptom they recognise in their own week.»"
    body = ("«Two or three sentences. Describe it the way they would describe "
            "it to a friend — the WhatsApp thread, the spreadsheet, the quote "
            "nobody followed up. Do not sell yet.»")
    note = "↳ «Optional margin note — a caveat or an aside.»"
    ts, tl = fit(title, TITLE_TTF, 44, CONTENT_W, 3)
    bs, bl = fit(body, BODY_TTF, 21, CONTENT_W, 6)
    ns, nl = fit(note, BODY_TTF, 15, CONTENT_W - 0.6, 3)
    stack([
        (0.42, lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.3, "The problem",
                                 color=pal["kicker"], **KICKER)),
        (block_h(tl, ts, 1.16) + 0.32,
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 2.0, title, font=TITLE_FONT,
                           size=ts, color=pal["text"], leading=1.16)),
        (block_h(bl, bs, 1.5) + 0.5,
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 2.4, body, font=BODY_FONT,
                           size=bs, color=pal["text"], leading=1.5)),
        (block_h(nl, ns, 1.45),
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.9, note, font=BODY_FONT,
                           size=ns, color=pal["muted"], leading=1.45)),
    ])


def s_stat(prs):
    pal = LIGHT
    s = new_slide(prs, pal, notes=(
        "STAT — one number, and it must be YOUR number.\n"
        "Every diagnostic claim on this account traces to something you "
        "observed: your audit, client work, logs. Never an imported statistic. "
        "The source line is not optional — if you cannot fill it, cut the "
        "slide.\nThe figure is this slide's amber. Nothing else on it is."))
    label, value, unit = "«What the number measures»", "«NN»", "«%»"
    caption = "«What that number costs them, in one sentence.»"
    source = ("Source: «your own audit, client log or observation — never an "
              "imported statistic»")
    cs, cl = fit(caption, BODY_TTF, 24, CONTENT_W, 3)
    ss, sl = fit(source, BODY_TTF, 13, CONTENT_W - 1.4, 3)

    def figure(y):
        box = textbox(s, MARGIN, y, CONTENT_W, 1.9, value, font=TITLE_FONT,
                      size=118, color=pal["amber"], leading=1.0)
        run = box.text_frame.paragraphs[0].add_run()
        run.text = " " + unit
        run.font.name = TITLE_FONT
        run.font.size = Pt(52)
        run.font.color.rgb = rgb(pal["amber"])

    def src(y):
        hairline(s, MARGIN, y - 0.22, CONTENT_W * 0.62, pal["line"])
        textbox(s, MARGIN, y, CONTENT_W - 1.4, 0.9, source, font=BODY_FONT,
                size=ss, color=pal["muted"], leading=1.45)

    stack([
        (0.46, lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.3, label,
                                 color=pal["kicker"], **KICKER)),
        (1.9, figure),
        (block_h(cl, cs, 1.4) + 0.62, lambda y: textbox(
            s, MARGIN, y, CONTENT_W, 1.3, caption, font=BODY_FONT, size=cs,
            color=pal["text"], leading=1.4)),
        (block_h(sl, ss, 1.45), src),
    ])


def s_steps(prs):
    pal = LIGHT
    s = new_slide(prs, pal, notes=(
        "THE FIX — the framework, named.\n"
        "Three to five steps, each an action in the imperative, short enough "
        "to read at a glance. Numerals stay teal-grey: no amber on this slide "
        "unless you drop the accent everywhere else."))
    title = "«The framework, named.»"
    items = ["«Step one — an action, in the imperative.»",
             "«Step two — short enough to read at a glance.»",
             "«Step three»",
             "«Step four — five items is the ceiling.»"]
    ts, tl = fit(title, TITLE_TTF, 44, CONTENT_W, 2)
    num_w, gap = 0.5, 0.22
    item_w = CONTENT_W - num_w - gap
    sized = [fit(i, BODY_TTF, 20, item_w, 3) for i in items]

    blocks = [
        (0.42, lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.3, "The fix",
                                 color=pal["kicker"], **KICKER)),
        (block_h(tl, ts, 1.16) + 0.42,
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 1.4, title, font=TITLE_FONT,
                           size=ts, color=pal["text"], leading=1.16)),
    ]
    for n, ((isz, ilines), text) in enumerate(zip(sized, items), 1):
        h = block_h(ilines, isz, 1.4) + 0.46

        def draw(y, n=n, text=text, isz=isz):
            hairline(s, MARGIN, y - 0.16, CONTENT_W, pal["line"])
            textbox(s, MARGIN, y + 0.03, num_w, 0.4, f"{n:02d}", font=BODY_FONT,
                    size=isz - 3, color=pal["muted"], tracking=1.2)
            textbox(s, MARGIN + num_w + gap, y, item_w, 1.2, text,
                    font=BODY_FONT, size=isz, color=pal["text"], leading=1.4)
        blocks.append((h, draw))
    blocks.append((0.2, lambda y: hairline(s, MARGIN, y - 0.16, CONTENT_W,
                                           pal["line"])))
    stack(blocks)


def s_compare(prs):
    pal = LIGHT
    s = new_slide(prs, pal, notes=(
        "BEFORE / AFTER — describe a process, not an adjective.\n"
        "If the 'after' is vaguer than the 'before', the contrast reads as "
        "marketing. The arrow is the amber."))
    title = "«Before and after, in one line.»"
    before = "«The current state. A process, not an adjective.»"
    after = ("«The state after the fix, just as concrete — or the contrast "
             "reads as marketing.»")
    ts, tl = fit(title, TITLE_TTF, 44, CONTENT_W, 2)
    pad = 0.34
    bs, bl = fit(before, BODY_TTF, 19, CONTENT_W - 2 * pad, 4)
    as_, al = fit(after, BODY_TTF, 19, CONTENT_W - 2 * pad, 4)
    bh = block_h(bl, bs, 1.42) + 1.0
    ah = block_h(al, as_, 1.42) + 1.0

    def pane(y, h, fill, label, text, size, tcol, lcol):
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(y),
                                 Inches(CONTENT_W), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(fill)
        box.line.fill.background()
        box.shadow.inherit = False
        textbox(s, MARGIN + pad, y + pad, CONTENT_W - 2 * pad, 0.3, label,
                font=BODY_FONT, size=12, color=lcol, tracking=2.2, caps=True)
        textbox(s, MARGIN + pad, y + pad + 0.42, CONTENT_W - 2 * pad, h - pad,
                text, font=BODY_FONT, size=size, color=tcol, leading=1.42)

    stack([
        (0.42, lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.3, "What changes",
                                 color=pal["kicker"], **KICKER)),
        (block_h(tl, ts, 1.16) + 0.40,
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 1.4, title, font=TITLE_FONT,
                           size=ts, color=pal["text"], leading=1.16)),
        (bh, lambda y: pane(y, bh, "EAE4D5", "Before", before, bs,
                            LIGHT["text"], LIGHT["muted"])),
        (0.62, lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.5, "↓",
                                 font=BODY_FONT, size=24, color=pal["amber"],
                                 align=PP_ALIGN.CENTER)),
        (ah, lambda y: pane(y, ah, DARK["bg"], "After", after, as_,
                            DARK["text"], DARK["muted"])),
    ])


def s_quote(prs):
    pal = LIGHT
    s = new_slide(prs, pal, notes=(
        "PULL QUOTE — the screenshot slide.\n"
        "Yours or a client's, attributed either way. If it is a client's, you "
        "need their permission before it goes out."))
    quote = "«A line worth screenshotting.»"
    attrib = "«Role and sector — or your own name»"
    qs, ql = fit(quote, TITLE_TTF, 52, CONTENT_W, 4)
    stack([
        (0.9, lambda y: textbox(s, MARGIN, y - 0.35, CONTENT_W, 1.0, "“",
                                font=TITLE_FONT, size=90, color="E8C98F",
                                leading=1.0)),
        (block_h(ql, qs, 1.22) + 0.55,
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 3.0, quote, font=TITLE_FONT,
                           size=qs, color=pal["text"], leading=1.22)),
        (0.4, lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.4, attrib,
                                color=pal["kicker"], **KICKER)),
    ])


def s_cta(prs):
    pal = DARK
    s = new_slide(prs, pal, notes=(
        "THE ASK — state it once.\n"
        "Be specific: vague asks get vague replies. Spell the DM keyword the "
        "same way here, in the caption, and in Aiden's rules, or the "
        "automation will not fire.\nThe keyword box is the amber."))
    title = "«The ask, stated once.»"
    line = ("«What they get, and what it costs them to ask. Be specific — "
            "vague asks get vague replies.»")
    keyword = "«DM KEYWORD»"
    note = "«What happens after they send it.»"
    ts, tl = fit(title, TITLE_TTF, 46, CONTENT_W, 2)
    ls, ll = fit(line, BODY_TTF, 21, CONTENT_W, 4)
    ns, nl = fit(note, BODY_TTF, 14, CONTENT_W, 3)

    def chip(y):
        w = 3.5
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(y),
                                 Inches(w), Inches(0.78))
        box.fill.background()
        box.line.color.rgb = rgb(pal["amber"])
        box.line.width = Pt(1.5)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = keyword
        run.font.name = BODY_FONT
        run.font.size = Pt(19)
        run.font.color.rgb = rgb(pal["amber"])
        run.font._rPr.set("spc", "220")

    stack([
        (0.42, lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.3, "«Next step»",
                                 color=pal["kicker"], **KICKER)),
        (block_h(tl, ts, 1.16) + 0.34,
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 1.6, title, font=TITLE_FONT,
                           size=ts, color=pal["text"], leading=1.16)),
        (block_h(ll, ls, 1.5) + 0.5,
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 1.8, line, font=BODY_FONT,
                           size=ls, color=pal["text"], leading=1.5)),
        (1.28, chip),
        (block_h(nl, ns, 1.45),
         lambda y: textbox(s, MARGIN, y, CONTENT_W, 0.8, note, font=BODY_FONT,
                           size=ns, color=pal["muted"], leading=1.45)),
    ])


def s_howto(prs):
    """Deliberately last, deliberately loud. Not part of any post."""
    pal = LIGHT
    s = new_slide(prs, pal, notes="Delete or skip this slide when you export.")
    rules = (
        "1.  Duplicate this deck for every post — File ▸ Make a copy.\n"
        "2.  Replace every «guillemet» slot. If one is left, do not post.\n"
        "3.  ONE amber element per slide. If two numbers are amber, one is wrong.\n"
        "4.  Every number carries a source line, and the source is your own\n"
        "     audit, client log or observation. Never an imported statistic.\n"
        "5.  Arabic goes to a native speaker before it goes out. Always.\n"
        "6.  Delete the slides you don't need. Six to eight is the range.\n"
        "7.  Export: File ▸ Download ▸ PNG for the current slide, one slide at\n"
        "     a time, or PDF for the set.\n"
        "8.  Fonts are Freehand (titles) and Tenor Sans (body). Tenor Sans\n"
        "     stands in for Optima, which Google Slides cannot serve."
    )
    textbox(s, MARGIN, 1.5, CONTENT_W, 0.5, "NOT PART OF THE POST",
            font=BODY_FONT, size=13, color="A6431F", tracking=2.4, caps=True)
    textbox(s, MARGIN, 2.1, CONTENT_W, 1.0, "How to use this template",
            font=TITLE_FONT, size=38, color=LIGHT["text"], leading=1.16)
    hairline(s, MARGIN, 3.15, CONTENT_W, LIGHT["line"])
    textbox(s, MARGIN, 3.45, CONTENT_W, 6.0, rules, font=BODY_FONT, size=16,
            color=LIGHT["text"], leading=1.55)


def main():
    for f in (TITLE_TTF, BODY_TTF):
        if not f.exists():
            raise SystemExit(f"missing font {f}. See README - they are fetched "
                             "from Google Fonts into tools/carousel/fonts/.")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(PAGE_W), Inches(PAGE_H)
    for build in (s_cover, s_text, s_stat, s_steps, s_compare, s_quote, s_cta,
                  s_howto):
        build(prs)
    OUT.mkdir(parents=True, exist_ok=True)
    dest = OUT / "AI-Profit-Lab-Carousel-Template.pptx"
    prs.save(dest)
    print(f"{len(prs.slides.__iter__.__self__._sldIdLst)} slides -> {dest}")
    return dest


if __name__ == "__main__":
    main()
